"""
051_master_batch_processor_cross.py
クロス集計 統合マスタースクリプト (個別出力・自動ファイル名対応版)

【特徴】
- 共通データ（回答CSV、マスタ、テンプレート）の読み込みは「最初の1回のみ」で爆速化
- 出力はバッチ条件のCSV名に合わせて動的に生成（例: _011_shisho の部分を自動抽出して結合）
- ファイルごとにWorkbookを破棄・gc.collectするためメモリリークの危険性ゼロ
- A列〜Z列の列幅を「10」で固定し、重い動的幅調整処理を排除してさらに高速化
"""

import pandas as pd
import numpy as np
import os
import sys
import hashlib
import unicodedata
import re
import datetime
import gc
import time
import traceback
import glob
import psutil

from openpyxl import load_workbook, Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment


class IntegratedCrossTabGenerator:
    SUMMARY_TEMPLATES = [
        "【分析要約】{cond_desc}（対象{total_ja}ＪＡ）を対象に「{target_name}」の回答傾向をクロス集計した結果、該当ＪＡ数が多い順に{ranking_str}となった。\n"
        "この結果から、同グループ内では上位項目に対する課題認識や取り組み方針が強く共有されていることが窺える。\n"
        "一方で下位にとどまった項目については、現場での実装ハードルやリソース不足が影響している可能性がある。今後の対策として、支所ごとの実施率ギャップを詳細に分析し、上位項目の成功事例を横展開しつつ、優先順位に応じた支援策やリソースの重点配分を行うことが有効である。",

        "【分析要約】{cond_desc}（計{total_ja}ＪＡ）の「{target_name}」に対する動向を確認した。最多回答から順に{ranking_str}と続いており、特定の施策や課題に対する関心の偏りが明確に表れている。\n"
        "上位に挙がった項目は、既に該当ＪＡ間で一定の有用性が認められていると推測される。逆に回答が少ない項目は、導入メリットの周知不足や体制整備の遅れが課題となっている可能性が高い。\n"
        "組織全体としての底上げを図るため、先行して取り組むＪＡの知見を集約し、マニュアル化や研修等を通じて積極的にノウハウを共有するアプローチが求められる。",

        "【分析要約】交差分析によれば、{cond_desc}（全{total_ja}ＪＡ）において、「{target_name}」は{ranking_str}の順で回答が集まる結果となった。\n"
        "この分布は、管内の農業構造や経営環境の違いが、各ＪＡの施策優先度に直接反映されていることを示唆している。特に上位項目は地域共通の急務である一方、下位項目は特定条件下のＪＡにのみ適合しているか、未着手の潜在的課題である可能性が考えられる。\n"
        "画一的な支援策ではなく、各ＪＡの実情に寄り添った伴走型のサポート体制を構築し、個別最適化された対応策を講じることが重要となる。"
    ]

    def __init__(self, ans_file, mapping_file, template_file, batch_files, output_dir):
        self.ans_file = ans_file
        self.mapping_file = mapping_file
        self.template_file = template_file
        self.batch_files = batch_files
        self.output_dir = output_dir

        self.df_merged = None
        self.col_map = None
        self.known_areas = None
        self.known_branches = None
        self.template_map = None
        self.sheet_name_counts = {}

    @staticmethod
    def read_csv_auto(filepath, **kwargs):
        kwargs.pop('encoding', None)
        try:
            return pd.read_csv(filepath, encoding='utf-8', **kwargs)
        except UnicodeDecodeError:
            return pd.read_csv(filepath, encoding='cp932', **kwargs)

    @staticmethod
    def normalize_ja_name(name):
        if pd.isna(name): return ""
        val = unicodedata.normalize('NFKC', str(name))
        val = re.sub(r'\s+', '', val)
        val = val.upper()
        val = val.replace('農業協同組合', 'JA').replace('農協', 'JA')
        return val

    @staticmethod
    def _clean_str(v):
        v_str = str(v).strip()
        return '' if v_str == 'nan' else v_str

    @staticmethod
    def parse_flag(flag_str):
        # 1 のフラグは無視する（元の仕様）
        f = str(flag_str).strip()
        if f in ['11', '１１', '1 1', '〇〇', '○○']: return 1, 1
        elif f in ['10', '１０', '〇×', '○×']: return 1, 0
        elif f in ['0', '０', '0.0', '×〇', '×○']: return 0, 1
        elif f in ['-1', '－１', '-1.0', '××']: return 0, 0
        else: return None, None

    @staticmethod
    def is_selected(val, target_code=None):
        if pd.isna(val): return False
        s_val = str(val).strip()
        if s_val in ['1', '1.0', 'True', 'TRUE', '〇', '○']: return True
        if target_code and s_val == str(target_code): return True
        return False

    @staticmethod
    def is_match_qid(col_qid, target_qid):
        col_qid_str = str(col_qid)
        target_qid_str = str(target_qid)
        return (col_qid_str == target_qid_str) or \
               col_qid_str.startswith(target_qid_str + '_') or \
               col_qid_str.startswith(target_qid_str + '-')

    def _safe_save(self, wb, output_path, retries=5, retry_delay=1.0):
        tmp_path = output_path + ".partial"
        wb.save(tmp_path)
        last_err = None
        for attempt in range(retries):
            try:
                os.replace(tmp_path, output_path)
                return
            except PermissionError as e:
                last_err = e
                if attempt < retries - 1:
                    time.sleep(retry_delay)
        raise last_err

    def load_common_data(self):
        print("⏳ 共通データ（回答データ・マスタ・テンプレート）をメモリに展開しています...")

        df_tmpl = self.read_csv_auto(self.template_file, header=None, dtype=str)
        def resolve_ref(ref_str):
            if pd.isna(ref_str) or not str(ref_str).strip(): return ""
            ref = str(ref_str).strip()
            match = re.match(r'^([A-Za-z]+)(\d+)$', ref)
            if not match: return ref
            col_str, row_str = match.groups()
            col_idx = 0
            for char in col_str.upper():
                col_idx = col_idx * 26 + (ord(char) - ord('A') + 1)
            col_idx -= 1
            row_idx = int(row_str) - 1
            try:
                val = df_tmpl.iloc[row_idx, col_idx]
                return str(val).strip() if not pd.isna(val) else ""
            except IndexError:
                return ""

        self.template_map = {}
        df_ids = df_tmpl.iloc[8:418, 7:11].copy()
        df_ids = df_ids.ffill()
        for r_idx in range(8, min(418, len(df_tmpl))):
            q_id = self._clean_str(df_ids.iloc[r_idx - 8, 1])
            m_id = self._clean_str(df_ids.iloc[r_idx - 8, 2])
            s_id = self._clean_str(df_ids.iloc[r_idx - 8, 3])
            if not q_id: continue

            self.template_map[(q_id, m_id, s_id)] = {
                'large_name': resolve_ref(df_tmpl.iloc[r_idx, 14]),
                'q_name': resolve_ref(df_tmpl.iloc[r_idx, 15]),
                'mid_name': resolve_ref(df_tmpl.iloc[r_idx, 16]),
                'sub_name': resolve_ref(df_tmpl.iloc[r_idx, 17])
            }

        df_headers = self.read_csv_auto(self.ans_file, skiprows=11, nrows=8, header=None, keep_default_na=False)
        self.col_map = {}
        for i in range(df_headers.shape[1]):
            self.col_map[i] = {
                'q_id': self._clean_str(df_headers.iloc[1, i] if df_headers.shape[0] > 1 else ""),
                'mid_item': self._clean_str(df_headers.iloc[2, i] if df_headers.shape[0] > 2 else ""),
                'sub_item': self._clean_str(df_headers.iloc[3, i] if df_headers.shape[0] > 3 else ""),
                'q_name': self._clean_str(df_headers.iloc[5, i] if df_headers.shape[0] > 5 else ""),
                'mid_name': self._clean_str(df_headers.iloc[6, i] if df_headers.shape[0] > 6 else ""),
                'sub_name': self._clean_str(df_headers.iloc[7, i] if df_headers.shape[0] > 7 else "")
            }

        df_data = self.read_csv_auto(self.ans_file, skiprows=19, header=None)
        df_data.replace({'True': True, 'False': False, 'TRUE': True, 'FALSE': False}, inplace=True)
        df_data.rename(columns={0: 'JA名'}, inplace=True)
        df_data['JA名'] = df_data['JA名'].astype(str).str.strip()
        df_data['__MatchKey__'] = df_data['JA名'].apply(self.normalize_ja_name)

        df_mapping = self.read_csv_auto(self.mapping_file, skiprows=7, usecols=[1, 2, 4], names=['エリア名', '支所名', 'JA名'])
        df_mapping.dropna(subset=['JA名'], inplace=True)
        df_mapping['JA名'] = df_mapping['JA名'].astype(str).str.strip()
        df_mapping['エリア名'] = df_mapping['エリア名'].astype(str).str.strip()
        df_mapping['支所名'] = df_mapping['支所名'].astype(str).str.strip()
        df_mapping['__MatchKey__'] = df_mapping['JA名'].apply(self.normalize_ja_name)
        df_mapping = df_mapping.drop_duplicates(subset='__MatchKey__', keep='first').drop(columns=['JA名'])

        self.df_merged = pd.merge(df_data, df_mapping, on='__MatchKey__', how='left')
        self.df_merged['支所名'] = self.df_merged['支所名'].fillna('不明')
        self.df_merged['エリア名'] = self.df_merged['エリア名'].fillna('不明')

        self.known_areas = set(self.df_merged['エリア名'].unique())
        self.known_branches = set(self.df_merged['支所名'].unique())

        print("✅ 共通データのロードが完了しました。")

    def _resolve_cols(self, qid, mid, sub):
        if not mid and not sub:
            if qid in self.known_areas: return [{'type': 'AREA', 'name': qid, 'label': qid, 'match_key': qid}]
            if qid in self.known_branches: return [{'type': 'BRANCH', 'name': qid, 'label': qid, 'match_key': qid}]

        cols = []
        for k, v in self.col_map.items():
            if self.is_match_qid(v['q_id'], qid) and v['mid_item'] == mid and (sub == '' or sub == v['sub_item']):
                display_label = v['sub_name'] or v['sub_item'] or v['mid_name'] or v['q_name'] or f"({v['q_id']}_項目なし)"
                cols.append({'type': 'QUESTION', 'key': k, 'label': display_label, 'match_key': v['sub_item'], 'q_id': v['q_id']})
        return cols

    def process_all_batches(self):
        start_time = datetime.datetime.now()
        
        self.load_common_data()

        total_batches = len(self.batch_files)
        total_generated_sheets = 0

        print(f"\n🚀 全 {total_batches} 個のバッチファイルの一括処理を開始します！")
        print(f"📁 出力方式: 各バッチごとに個別のExcelファイルを動的生成します")

        for b_idx, batch_file in enumerate(self.batch_files, 1):
            if not os.path.exists(batch_file):
                continue

            basename = os.path.basename(batch_file)
            
            match = re.search(r'\d{3}_batch_processing_cross(.*)\.csv$', basename)
            if match:
                suffix = match.group(1) 
            else:
                suffix = f"_batch_{b_idx}"
                
            out_filename = f"052_batch_analysis_master_cross{suffix}.xlsx"
            out_filepath = os.path.join(self.output_dir, out_filename)

            print(f"\n▶️ [{b_idx}/{total_batches}] 処理中: {basename}")
            print(f"  └─ 出力先: {out_filename}")
            
            try:
                is_new_output = not os.path.exists(out_filepath)
                try:
                    if not is_new_output:
                        wb = load_workbook(out_filepath)
                    else:
                        wb = Workbook()
                except Exception as e:
                    print(f"  ℹ️ 情報: 既存ファイルの読み込みに失敗したため、新規作成します（詳細: {e}）。")
                    wb = Workbook()
                    is_new_output = True

                if is_new_output and "Sheet" in wb.sheetnames:
                    wb.remove(wb["Sheet"])
                if "一括クロス集計結果" in wb.sheetnames:
                    wb.remove(wb["一括クロス集計結果"])

                self.sheet_name_counts = {}
                file_generated_sheets = 0

                df_batch = self.read_csv_auto(batch_file, header=None, dtype=str)
                matrix_rows, matrix_cols = df_batch.shape
                tasks = []
                unrecognized_flags = set()

                for r in range(3, matrix_rows):
                    valid_target_cols = []
                    for c in range(3, matrix_cols):
                        raw_flag = str(df_batch.iloc[r, c]).strip()
                        if not raw_flag or raw_flag == 'nan': continue
                        cond_val, targ_val = self.parse_flag(raw_flag)
                        if cond_val is not None:
                            valid_target_cols.append((c, raw_flag, cond_val, targ_val))
                        else:
                            unrecognized_flags.add(raw_flag)
                    if valid_target_cols:
                        tasks.append((r, valid_target_cols))

                if unrecognized_flags:
                    print(f"  ⚠️ 警告: 未定義のフラグ値が見つかり、スキップされました: {unrecognized_flags}")

                if not tasks:
                    print(f"  ⚠️ 警告: 有効なクロス集計対象が1件もありませんでした。")
                else:
                    for task_idx, (r, valid_target_cols) in enumerate(tasks, 1):
                        cond_qid = self._clean_str(df_batch.iloc[r, 0])
                        cond_mid = self._clean_str(df_batch.iloc[r, 1])
                        cond_sub = self._clean_str(df_batch.iloc[r, 2])

                        cond_cols_list = self._resolve_cols(cond_qid, cond_mid, cond_sub)
                        if not cond_cols_list: continue

                        cond_type = cond_cols_list[0]['type']

                        raw_sheet_name = f"{cond_qid}_{cond_sub}" if cond_sub else f"{cond_qid}"
                        sanitized_name = re.sub(r'[\\*?:/\[\]]', '_', raw_sheet_name).strip("'")
                        if not sanitized_name:
                            sanitized_name = "無題"
                        if len(sanitized_name) > 28:
                            name_hash = hashlib.md5(sanitized_name.encode("utf-8")).hexdigest()[:4]
                            base_sheet_name = sanitized_name[:23] + "_" + name_hash
                        else:
                            base_sheet_name = sanitized_name

                        if base_sheet_name not in self.sheet_name_counts:
                            self.sheet_name_counts[base_sheet_name] = 1
                            sheet_name = base_sheet_name
                        else:
                            sheet_name = f"{base_sheet_name}_{self.sheet_name_counts[base_sheet_name]}"
                            self.sheet_name_counts[base_sheet_name] += 1

                        row1 = ["", "", "", ""]
                        row2 = ["", "", "", ""]
                        row3 = ["設問ID", "中項目", "小項目", "総数"]

                        expanded_targets = []
                        for c, raw_flag, cond_val, targ_val in valid_target_cols:
                            t_qid = self._clean_str(df_batch.iloc[0, c])
                            t_mid = self._clean_str(df_batch.iloc[1, c])
                            t_sub = self._clean_str(df_batch.iloc[2, c])
                            t_cols = self._resolve_cols(t_qid, t_mid, t_sub)

                            for tc in t_cols:
                                expanded_targets.append({'tc': tc, 'cond_val': cond_val, 'targ_val': targ_val})
                                row1.append(tc.get('q_id', tc.get('name', '')))
                                row2.append(t_mid)
                                row3.append(tc.get('label', tc.get('name', '')))

                        summary_data = []
                        for cc in cond_cols_list:
                            row_data = [
                                cc.get('q_id', cond_qid), cond_mid,
                                cc.get('label', cond_sub) if cc.get('label') else cc.get('name', ''), 0
                            ]

                            if cc['type'] == 'AREA': cm = (self.df_merged['エリア名'] == cc['name']).values
                            elif cc['type'] == 'BRANCH': cm = (self.df_merged['支所名'] == cc['name']).values
                            else: cm = self.df_merged[cc['key']].apply(lambda x: self.is_selected(x, cc['match_key'])).values

                            row_data[3] = cm.sum()

                            for et in expanded_targets:
                                tc = et['tc']
                                current_cm = cm if et['cond_val'] == 1 else ~cm
                                if tc['type'] == 'AREA': tm = (self.df_merged['エリア名'] == tc['name']).values
                                elif tc['type'] == 'BRANCH': tm = (self.df_merged['支所名'] == tc['name']).values
                                else: tm = self.df_merged[tc['key']].apply(lambda x: self.is_selected(x, tc['match_key'])).values

                                current_tm = tm if et['targ_val'] == 1 else ~tm
                                row_data.append((current_cm & current_tm).sum())

                            summary_data.append(row_data)

                        if sheet_name in wb.sheetnames:
                            wb.remove(wb[sheet_name])
                        ws = wb.create_sheet(sheet_name)

                        ws.cell(row=1, column=1, value=f"▼ クロス集計サマリー ({sheet_name})").font = Font(bold=True, size=12)
                        for c_idx, val in enumerate(row1, 1): ws.cell(row=2, column=c_idx, value=val).font = Font(color="008000")
                        for c_idx, val in enumerate(row2, 1): ws.cell(row=3, column=c_idx, value=val).font = Font(color="008000")
                        for c_idx, val in enumerate(row3, 1):
                            cell = ws.cell(row=4, column=c_idx, value=val)
                            cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
                            cell.font = Font(bold=True)

                        for r_i, row_data in enumerate(summary_data, 5):
                            for c_idx, val in enumerate(row_data, 1):
                                ws.cell(row=r_i, column=c_idx, value=val)

                        global_row = max(82, len(summary_data) + 6)

                        for c, raw_flag, cond_val, targ_val in valid_target_cols:
                            target_qid = self._clean_str(df_batch.iloc[0, c])
                            target_mid = self._clean_str(df_batch.iloc[1, c])
                            target_sub = self._clean_str(df_batch.iloc[2, c])
                            target_cols_list = self._resolve_cols(target_qid, target_mid, target_sub)
                            if not target_cols_list: continue
                            target_type = target_cols_list[0]['type']

                            if cond_type == 'AREA': t_cond = {'large_name': '地域属性', 'q_name': 'エリア指定', 'mid_name': '', 'sub_name': cond_qid}
                            elif cond_type == 'BRANCH': t_cond = {'large_name': '地域属性', 'q_name': '支所指定', 'mid_name': '', 'sub_name': cond_qid}
                            else: t_cond = self.template_map.get((cond_qid, cond_mid, cond_sub), {})

                            if target_type in ['AREA', 'BRANCH']: t_targ = {'large_name': '地域属性', 'q_name': '地域指定', 'mid_name': '', 'sub_name': target_qid}
                            else: t_targ = self.template_map.get((target_qid, target_mid, target_sub), {})

                            combo_labels = []
                            combo_map = {}
                            combo_arrays = {}
                            cond_masks = []

                            for i, cc in enumerate(cond_cols_list):
                                if cond_type == 'AREA': cm = (self.df_merged['エリア名'].values == cc['name'])
                                elif cond_type == 'BRANCH': cm = (self.df_merged['支所名'].values == cc['name'])
                                else: cm = self.df_merged[cc['key']].apply(lambda x: self.is_selected(x, cc['match_key'])).values
                                if cond_val == 0: cm = ~cm
                                cond_masks.append(cm)

                            cond_only_mask = np.zeros(len(self.df_merged), dtype=bool)
                            for cm in cond_masks: cond_only_mask |= cm

                            targ_masks = []
                            for j, tc in enumerate(target_cols_list):
                                if target_type == 'AREA': tm = (self.df_merged['エリア名'].values == tc['name'])
                                elif target_type == 'BRANCH': tm = (self.df_merged['支所名'].values == tc['name'])
                                else: tm = self.df_merged[tc['key']].apply(lambda x: self.is_selected(x, tc['match_key'])).values
                                if targ_val == 0: tm = ~tm
                                targ_masks.append(tm)

                            for i, cc in enumerate(cond_cols_list):
                                for j, tc in enumerate(target_cols_list):
                                    col_name_base = f"[{cc['label']}] {tc['label']}" if len(cond_cols_list) > 1 else tc['label']
                                    col_name = col_name_base
                                    suffix_idx = 1
                                    while col_name in combo_labels:
                                        col_name = f"{col_name_base}_{suffix_idx}"
                                        suffix_idx += 1
                                    combo_labels.append(col_name)
                                    combo_map[(i, j)] = col_name
                                    combo_arrays[col_name] = cond_masks[i] & targ_masks[j]

                            total_ja = int(cond_only_mask.sum())

                            transposed_data = [[""] + [c['label'] for c in cond_cols_list]]
                            for j, tc in enumerate(target_cols_list):
                                row = [tc['label']]
                                for i, cc in enumerate(cond_cols_list):
                                    row.append(int(combo_arrays[combo_map[(i, j)]].sum()))
                                transposed_data.append(row)

                            detail_dict = {
                                'エリア名': self.df_merged['エリア名'].values[cond_only_mask],
                                '支所名': self.df_merged['支所名'].values[cond_only_mask],
                                'JA名': self.df_merged['JA名'].values[cond_only_mask],
                            }
                            for col_name in combo_labels:
                                detail_dict[col_name] = np.where(combo_arrays[col_name][cond_only_mask], '〇', '')

                            df_detail = pd.DataFrame(detail_dict, columns=['エリア名', '支所名', 'JA名'] + combo_labels)
                            df_detail = df_detail.sort_values(['エリア名', '支所名'])

                            summary_series = pd.Series({col: int(arr.sum()) for col, arr in combo_arrays.items()}).sort_values(ascending=False)
                            ranking_str = "、次いで".join([f"（{k}）{v}JA" for k, v in summary_series.items()][:3]) if not summary_series.empty and total_ja > 0 else "該当回答なし"

                            cond_desc = f"「{cond_qid}」{'を選択しなかった' if cond_val==0 else 'を選択した'}ＪＡ"
                            targ_desc = f"{target_qid}{'（未選択）' if targ_val==0 else '（選択）'}"

                            hash_val = int(hashlib.md5(f"{cond_qid}_x_{target_qid}".encode('utf-8')).hexdigest(), 16)
                            a1_text = self.SUMMARY_TEMPLATES[hash_val % len(self.SUMMARY_TEMPLATES)].format(
                                cond_desc=cond_desc, target_name=targ_desc if target_type == 'QUESTION' else "地域属性", total_ja=total_ja, ranking_str=ranking_str
                            )

                            base_row = global_row
                            bar_width = max(14, len(combo_labels) + 3)
                            for col_idx in range(1, bar_width + 1):
                                ws.cell(row=base_row, column=col_idx).fill = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
                            base_row += 2

                            ws.cell(row=base_row, column=1, value=f"【縦軸 (条件)】 大項目: {t_cond.get('large_name', '')}").font = Font(bold=True, color="0000FF")
                            ws.cell(row=base_row+1, column=1, value=f"【縦軸 (条件)】 設問名: {cond_qid} {t_cond.get('q_name', '')}").font = Font(bold=True, color="0000FF")
                            ws.cell(row=base_row+2, column=1, value=f"【縦軸 (条件)】 中項目: {t_cond.get('mid_name', '')}").font = Font(bold=True, color="0000FF")
                            ws.cell(row=base_row+3, column=1, value=f"【縦軸 (条件)】 小項目: {t_cond.get('sub_name', '一括')}").font = Font(bold=True, color="0000FF")

                            ws.cell(row=base_row+4, column=1, value=f"【横軸 (ターゲット)】 大項目: {t_targ.get('large_name', '')}").font = Font(bold=True, color="008000")
                            ws.cell(row=base_row+5, column=1, value=f"【横軸 (ターゲット)】 設問名: {target_qid} {t_targ.get('q_name', '')}").font = Font(bold=True, color="008000")
                            ws.cell(row=base_row+6, column=1, value=f"【横軸 (ターゲット)】 中項目: {t_targ.get('mid_name', '')}").font = Font(bold=True, color="008000")
                            ws.cell(row=base_row+7, column=1, value=f"【横軸 (ターゲット)】 小項目: {t_targ.get('sub_name', '一括')}").font = Font(bold=True, color="008000")

                            ws.cell(row=base_row+8, column=1, value=a1_text).font = Font(bold=False, size=10)
                            ws.merge_cells(start_row=base_row+8, start_column=1, end_row=base_row+8, end_column=8)
                            ws.cell(row=base_row+8, column=1).alignment = Alignment(wrap_text=True, vertical="top")
                            ws.row_dimensions[base_row+8].height = 80

                            curr_row = base_row + 9
                            ws.cell(row=curr_row, column=1, value=f"▼ クロス集計表").font = Font(bold=True, size=11)
                            curr_row += 1
                            for r_i, row_data in enumerate(transposed_data, curr_row):
                                for c_idx, val in enumerate(row_data, 1): ws.cell(row=r_i, column=c_idx, value=val)
                                curr_row = r_i

                            curr_row += 2
                            table3_data_start = curr_row + 1
                            ws.cell(row=curr_row, column=1, value=f"▼ 該当JA（{total_ja}件）の回答詳細").font = Font(bold=True)
                            for r_i, row in enumerate(dataframe_to_rows(df_detail, index=False, header=True), table3_data_start):
                                for c_idx, val in enumerate(row, 1):
                                    cell = ws.cell(row=r_i, column=c_idx, value=val)
                                    if r_i == table3_data_start:
                                        cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
                                        cell.font = Font(bold=True)

                            global_row = table3_data_start + len(df_detail) + 5
                            del cond_masks, targ_masks, combo_arrays, detail_dict, df_detail, transposed_data

                        # 💡 変更後: A列からZ列までの列幅を「10」で統一
                        for col_idx in range(1, 27):
                            col_letter = get_column_letter(col_idx)
                            ws.column_dimensions[col_letter].width = 10

                        total_generated_sheets += 1
                        file_generated_sheets += 1

                # 1ファイル分の処理が終わったら即座に保存してメモリから破棄
                if file_generated_sheets > 0 or not tasks:
                    print(f"  💾 {out_filename} を保存しています...")
                    try:
                        self._safe_save(wb, out_filepath)
                    except PermissionError:
                        alt_file = f"{out_filepath.replace('.xlsx', '')}_退避用_{datetime.datetime.now().strftime('%H%M%S')}.xlsx"
                        wb.save(alt_file)
                        print(f"  ⚠️ 警告: ファイルが開かれているため、退避用 {os.path.basename(alt_file)} に保存しました。")
                        
            except Exception:
                print(f"  ❌ 【異常終了】{basename} の処理中に予期しないエラーが発生しました。このバッチはスキップして続行します。")
                traceback.print_exc()
            finally:
                if 'wb' in locals():
                    try:
                        wb.close()
                    except Exception:
                        pass
                    del wb
                
                gc.collect()
                try:
                    rss_mb = psutil.Process(os.getpid()).memory_info().rss / 1024**2
                    print(f"  📊 完了後のメモリ使用量: {rss_mb:.0f} MB")
                except Exception:
                    pass

        elapsed = datetime.datetime.now() - start_time
        print(f"\n✨ 全処理が正常終了しました。(計 {total_batches} バッチファイル / {total_generated_sheets} シート作成完了)")
        print(f"⏱️ 統合マスター 総処理時間: {int(elapsed.total_seconds() // 60)}分 {int(elapsed.total_seconds() % 60)}秒")


if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.abspath(__file__))

    # フォルダ内にある 全てのバッチ用CSV を自動で探索する (* を活用)
    search_pattern = os.path.join(base_dir, "*batch_processing_cross_*.csv")
    batch_files_list = sorted(glob.glob(search_pattern))

    if not batch_files_list:
        print(f"❌ エラー: {base_dir} 内にバッチファイルが見つかりません。")
        sys.exit(1)

    # コマンドライン引数でバッチファイル単位の再開に対応
    if len(sys.argv) > 1:
        needle = sys.argv[1].lower()
        matched = [i for i, p in enumerate(batch_files_list) if needle in os.path.basename(p).lower()]
        if matched:
            start_idx = matched[0]
            print(f"⏩ 指定により {os.path.basename(batch_files_list[start_idx])} から処理を再開します。")
            batch_files_list = batch_files_list[start_idx:]
        else:
            print(f"❌ 無効な指定です: {sys.argv[1]}")
            sys.exit(1)

    generator = IntegratedCrossTabGenerator(
        ans_file=os.path.join(base_dir, '051_horizon_conv_cross_100ja_cleaned.csv'), # ←クリーンデータを読み込む
        mapping_file=os.path.join(base_dir, '081_question_id_mapping_real_cross.csv'),
        template_file=os.path.join(base_dir, '080_questionnaire_template_real_cross.csv'),
        batch_files=batch_files_list,
        output_dir=base_dir
    )
    generator.process_all_batches()