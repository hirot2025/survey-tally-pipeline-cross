"""
survey-tally-pipeline-simple
Copyright (c) 2026 HIROI Takayuki
Released under the MIT license
https://github.com/hirot2025/survey-tally-pipeline-simple/blob/main/LICENSE.md
"""

import os
import re
import pandas as pd
import openpyxl
from pathlib import Path

# ==========================================
# 設定
# ==========================================
TARGET_DIR = "000_ja100_cross"         # アンケートExcelが格納されているフォルダ
OUTPUT_DIR = "010_ja100_csv_cross"    # CSVを出力するフォルダ
SHEET_NAME = "集計用"

def sanitize_filename(name):
    """ファイル名として使用できない文字をアンダースコアに置換"""
    return re.sub(r'[\\/*?:"<>|]', '_', str(name))

def main():
    if not os.path.exists(TARGET_DIR):
        print(f"エラー: フォルダ '{TARGET_DIR}' が見つかりません。")
        return
        
    # 出力先フォルダの作成
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    excel_files = list(Path(TARGET_DIR).glob("*.xlsx"))
    
    if not excel_files:
        print(f"'{TARGET_DIR}' フォルダ内にExcelファイルが見つかりません。")
        return
        
    success_count = 0
    
    for file_path in excel_files:
        # Excelの一時ファイル（~$）はスキップ
        if file_path.name.startswith("~$"):
            continue
            
        try:
            # 1. openpyxlで B4セル（集計用ID）を取得
            wb = openpyxl.load_workbook(file_path, data_only=True)
            if SHEET_NAME not in wb.sheetnames:
                print(f"スキップ ({file_path.name}): '{SHEET_NAME}' シートがありません。")
                wb.close()
                continue
                
            ws = wb[SHEET_NAME]
            raw_id = ws["B4"].value
            wb.close() # メモリ解放のために閉じる
            
            # IDが空、または数式エラー(#REF!)の場合は元のファイル名をベースにする安全策
            if raw_id is None or str(raw_id) == "#REF!" or str(raw_id).strip() == "":
                file_id = f"{file_path.stem}_ID不明"
            else:
                file_id = sanitize_filename(str(raw_id).strip())
                
            out_name = f"{file_id}.csv"
            out_path = Path(OUTPUT_DIR) / out_name
            
            # 2. pandasでシート全体を読み込み、そのままCSVとして書き出し（マッピング変換などは行わない）
            df = pd.read_excel(file_path, sheet_name=SHEET_NAME, header=None)
            df.to_csv(out_path, index=False, header=False, encoding="utf-8-sig")
            
            print(f"出力完了: {out_name} (元ファイル: {file_path.name})")
            success_count += 1
            
        except Exception as e:
            print(f"エラー ({file_path.name}): {e}")
            
    print(f"\n処理完了: {success_count} 件のCSVファイルを '{OUTPUT_DIR}' に生成しました。")

if __name__ == "__main__":
    main()